"""Tests for the DOCX formatter module (in-memory)."""

from docx.enum.text import WD_LINE_SPACING  # type: ignore[import-untyped]
from docx.shared import Pt  # type: ignore[import-untyped]

from transcript_etl_pipeline.document.formatting_rules import BODY_FONT, LABEL_FONT, SPACING_RULES
from transcript_etl_pipeline.document.model import Label, Paragraph, SectionType
from transcript_etl_pipeline.formatters.docx_formatter import (
    _apply_font,  # pyright: ignore[reportPrivateUsage]
    _apply_spacing,  # pyright: ignore[reportPrivateUsage]
    _format_paragraph,  # pyright: ignore[reportPrivateUsage]
)


class FakeFont:
    """Minimal font stub to capture style assignments."""

    def __init__(self) -> None:
        self.name: str | None = None
        self.size: object | None = None
        self.bold: bool | None = None


class FakeDocxRun:
    """Minimal run stub for docx formatting."""

    def __init__(self, text: str = "") -> None:
        self.text = text
        self.font = FakeFont()


class FakeParagraphFormat:
    """Minimal paragraph format stub."""

    def __init__(self) -> None:
        self.space_before: object | None = None
        self.space_after: object | None = None
        self.line_spacing_rule: object | None = None


class FakeDocxParagraph:
    """Minimal paragraph stub for docx formatting."""

    def __init__(self, style: str | None = None) -> None:
        self.style = style
        self.runs: list[FakeDocxRun] = []
        self.paragraph_format = FakeParagraphFormat()

    def add_run(self, text: str) -> FakeDocxRun:
        run = FakeDocxRun(text)
        self.runs.append(run)
        return run

    @property
    def text(self) -> str:
        return "".join(run.text for run in self.runs)


class FakeDocxDocument:
    """Minimal document stub for docx formatting."""

    def __init__(self) -> None:
        self.paragraphs: list[FakeDocxParagraph] = []

    def add_paragraph(self, text: str = "", style: str | None = None) -> FakeDocxParagraph:
        paragraph = FakeDocxParagraph(style=style)
        if text:
            paragraph.add_run(text)
        self.paragraphs.append(paragraph)
        return paragraph


class TestDocxFormatterInMemory:
    """Tests for DOCX formatter helpers."""

    def test_format_paragraph_notes_header_heading_level(self) -> None:
        """Notes headers use the correct heading style."""
        docx_doc = FakeDocxDocument()
        paragraph = Paragraph(text="Agenda", heading_level=3)

        _format_paragraph(docx_doc, paragraph, SectionType.NOTES_HEADER)

        assert docx_doc.paragraphs[0].style == "Heading 3"
        assert docx_doc.paragraphs[0].text == "Agenda"

    def test_format_paragraph_notes_body_bullet_level_two(self) -> None:
        """Notes bullets use List Bullet 2 for level-two bullets."""
        docx_doc = FakeDocxDocument()
        paragraph = Paragraph(text="Nested", is_bullet=True, bullet_level=2)

        _format_paragraph(docx_doc, paragraph, SectionType.NOTES_BODY)

        assert docx_doc.paragraphs[0].style == "List Bullet 2"
        assert docx_doc.paragraphs[0].text == "Nested"

    def test_apply_spacing_sets_before_after_and_single_spacing(self) -> None:
        """Spacing helper applies before/after spacing and line spacing."""
        paragraph = FakeDocxParagraph()
        spacing = SPACING_RULES[SectionType.SPEAKER_PARAGRAPH]

        _apply_spacing(paragraph, spacing)

        assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE
        assert paragraph.paragraph_format.space_before == Pt(spacing.before_pt)
        assert paragraph.paragraph_format.space_after == Pt(spacing.after_pt)

    def test_format_paragraph_speaker_label_and_body(self) -> None:
        """Speaker paragraphs render bold label and normal body text."""
        docx_doc = FakeDocxDocument()
        paragraph = Paragraph(label=Label("Speaker:"), text="Hello world.")

        _format_paragraph(docx_doc, paragraph, SectionType.SPEAKER_PARAGRAPH)

        formatted = docx_doc.paragraphs[0]
        assert len(formatted.runs) == 2
        assert formatted.runs[0].font.bold is True
        assert formatted.runs[1].font.bold is False

    def test_apply_font_sets_font_style(self) -> None:
        """Font helper assigns the expected font properties."""
        run = FakeDocxRun("Text")

        _apply_font(run, BODY_FONT)

        assert run.font.name == BODY_FONT.name
        assert run.font.size == Pt(BODY_FONT.size_pt)
        assert run.font.bold == BODY_FONT.bold

        _apply_font(run, LABEL_FONT)
        assert run.font.bold == LABEL_FONT.bold
