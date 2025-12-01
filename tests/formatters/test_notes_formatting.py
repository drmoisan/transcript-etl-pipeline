"""Tests for notes formatting in all formatters.

Tests that notes sections (headers and body with bullets) are
properly formatted in DOCX, MD, and RTF output.
"""

from pathlib import Path

from docx import Document as DocxDocument  # type: ignore[import-untyped]

from transcript_etl_pipeline.document.model import (
    Document,
    DocumentSection,
    Paragraph,
    SectionType,
)
from transcript_etl_pipeline.formatters.docx_formatter import format_to_docx
from transcript_etl_pipeline.formatters.md_formatter import format_to_md
from transcript_etl_pipeline.formatters.rtf_formatter import format_to_rtf


class TestNotesFormattingDocx:
    """Tests for notes formatting in DOCX output."""

    def test_notes_header_uses_heading_style(self, tmp_path: Path) -> None:
        """Notes header paragraph uses Heading 2 Word style by default."""
        doc = Document()
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_HEADER,
                paragraphs=[Paragraph(text="Notes – 2025-01-01", heading_level=2)],
            )
        )

        output_path = tmp_path / "test.docx"
        format_to_docx(doc, str(output_path))

        # Read back and verify
        docx_doc = DocxDocument(str(output_path))  # type: ignore[no-untyped-call]
        para = docx_doc.paragraphs[0]
        assert para.text == "Notes – 2025-01-01"
        # Check that the style is Heading 2
        assert para.style.name == "Heading 2"

    def test_notes_header_h1_uses_heading1_style(self, tmp_path: Path) -> None:
        """Notes header with heading_level=1 uses Heading 1 Word style."""
        doc = Document()
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_HEADER,
                paragraphs=[Paragraph(text="Document Title", heading_level=1)],
            )
        )

        output_path = tmp_path / "test.docx"
        format_to_docx(doc, str(output_path))

        docx_doc = DocxDocument(str(output_path))  # type: ignore[no-untyped-call]
        assert docx_doc.paragraphs[0].style.name == "Heading 1"

    def test_notes_body_with_bullets_uses_list_style(self, tmp_path: Path) -> None:
        """Notes body with bullet paragraphs uses List Bullet style."""
        doc = Document()
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_BODY,
                paragraphs=[
                    Paragraph(text="First bullet", is_bullet=True, bullet_level=1),
                    Paragraph(text="Nested bullet", is_bullet=True, bullet_level=2),
                    Paragraph(text="Regular text", is_bullet=False),
                ],
            )
        )

        output_path = tmp_path / "test.docx"
        format_to_docx(doc, str(output_path))

        # Read back and verify
        docx_doc = DocxDocument(str(output_path))  # type: ignore[no-untyped-call]
        paras = docx_doc.paragraphs

        assert len(paras) == 3
        # First should use List Bullet style
        assert paras[0].style.name == "List Bullet"
        # Second should use List Bullet 2 style
        assert paras[1].style.name == "List Bullet 2"
        # Third should be regular paragraph


class TestNotesFormattingMd:
    """Tests for notes formatting in Markdown output."""

    def test_notes_header_is_h1(self, tmp_path: Path) -> None:
        """Notes header is formatted as H1 heading."""
        doc = Document()
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_HEADER,
                paragraphs=[Paragraph(text="Notes – 2025-01-01")],
            )
        )

        output_path = tmp_path / "test.md"
        format_to_md(doc, str(output_path))

        content = output_path.read_text()
        assert "# Notes – 2025-01-01" in content

    def test_notes_body_bullets_use_dash_prefix(self, tmp_path: Path) -> None:
        """Notes body bullets are prefixed with dash."""
        doc = Document()
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_BODY,
                paragraphs=[
                    Paragraph(text="First bullet", is_bullet=True),
                    Paragraph(text="Second bullet", is_bullet=True),
                ],
            )
        )

        output_path = tmp_path / "test.md"
        format_to_md(doc, str(output_path))

        content = output_path.read_text()
        assert "- First bullet" in content
        assert "- Second bullet" in content

    def test_notes_body_regular_text_no_prefix(self, tmp_path: Path) -> None:
        """Notes body regular text has no bullet prefix."""
        doc = Document()
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_BODY,
                paragraphs=[Paragraph(text="Regular paragraph", is_bullet=False)],
            )
        )

        output_path = tmp_path / "test.md"
        format_to_md(doc, str(output_path))

        content = output_path.read_text()
        assert "Regular paragraph" in content
        assert "- Regular paragraph" not in content


class TestNotesFormattingRtf:
    """Tests for notes formatting in RTF output."""

    def test_notes_header_is_bold(self, tmp_path: Path) -> None:
        """Notes header is formatted as bold in RTF."""
        doc = Document()
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_HEADER,
                paragraphs=[Paragraph(text="Notes – 2025-01-01")],
            )
        )

        output_path = tmp_path / "test.rtf"
        format_to_rtf(doc, str(output_path))

        content = output_path.read_text()
        # Check for bold tag around text
        assert r"{\b Notes" in content

    def test_notes_header_has_larger_font(self, tmp_path: Path) -> None:
        """Notes header uses larger font size (14pt = fs28)."""
        doc = Document()
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_HEADER,
                paragraphs=[Paragraph(text="Notes")],
            )
        )

        output_path = tmp_path / "test.rtf"
        format_to_rtf(doc, str(output_path))

        content = output_path.read_text()
        # fs28 = 14pt * 2 (RTF uses half-points)
        assert r"\fs28" in content

    def test_notes_body_bullets_have_bullet_char(self, tmp_path: Path) -> None:
        """Notes body bullets include bullet character."""
        doc = Document()
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_BODY,
                paragraphs=[Paragraph(text="Bullet point", is_bullet=True)],
            )
        )

        output_path = tmp_path / "test.rtf"
        format_to_rtf(doc, str(output_path))

        content = output_path.read_text()
        # Unicode bullet character code
        assert r"\u8226" in content


class TestFullDocumentWithNotes:
    """Tests for complete documents with notes sections."""

    def test_md_full_document_structure(self, tmp_path: Path) -> None:
        """Complete document with metadata, notes, and transcript."""
        doc = Document()
        doc.add_section(
            DocumentSection(
                section_type=SectionType.METADATA,
                paragraphs=[
                    Paragraph(text="Meeting: Team Standup"),
                    Paragraph(text="Date: 2025-01-01"),
                ],
            )
        )
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_HEADER,
                paragraphs=[Paragraph(text="Notes – 2025-01-01")],
            )
        )
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_BODY,
                paragraphs=[
                    Paragraph(text="Action item one", is_bullet=True),
                    Paragraph(text="Action item two", is_bullet=True),
                ],
            )
        )

        output_path = tmp_path / "test.md"
        format_to_md(doc, str(output_path))

        content = output_path.read_text()

        # Verify structure
        assert "Meeting: Team Standup" in content
        assert "# Notes – 2025-01-01" in content
        assert "- Action item one" in content
        assert "- Action item two" in content

    def test_docx_full_document_paragraph_count(self, tmp_path: Path) -> None:
        """Complete document has correct number of paragraphs."""
        doc = Document()
        doc.add_section(
            DocumentSection(
                section_type=SectionType.METADATA,
                paragraphs=[Paragraph(text="Meeting: Test")],
            )
        )
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_HEADER,
                paragraphs=[Paragraph(text="Notes", heading_level=2)],
            )
        )
        doc.add_section(
            DocumentSection(
                section_type=SectionType.NOTES_BODY,
                paragraphs=[
                    Paragraph(text="Bullet 1", is_bullet=True),
                    Paragraph(text="Bullet 2", is_bullet=True),
                ],
            )
        )

        output_path = tmp_path / "test.docx"
        format_to_docx(doc, str(output_path))

        docx_doc = DocxDocument(str(output_path))  # type: ignore[no-untyped-call]
        # Should have: 1 metadata + 1 header + 2 body = 4 paragraphs
        assert len(docx_doc.paragraphs) == 4
