"""End-to-end CLI tests for notes workflows.

Tests the CLI interface with notes processing to verify:
- Notes-only document creation
- Notes + transcript combined workflows
- Update mode (add/replace notes and transcript)
- Coverage of cli.py, document/reader.py, transform/notes.py
"""

from pathlib import Path

from docx import Document as DocxDocument  # type: ignore[import-untyped]

from tests.fixtures.multi_speaker import SPACEX_DISCUSSION
from transcript_etl_pipeline.cli import main


class TestCLINotesOnly:
    """Test CLI with notes-only workflows."""

    def test_notes_only_markdown(self, tmp_path: Path) -> None:
        """Test CLI creating Markdown document with notes only.

        Exercises: cli.main(), transform_notes(), format_to_md()
        """
        # Arrange: Create notes file
        notes_file = tmp_path / "notes.md"
        notes_file.write_text(
            "- Action item: Review Q4 results\n"
            "- Decision: Approve budget increase\n"
            "- Next meeting: January 15th\n",
            encoding="utf-8",
        )

        output_file = tmp_path / "notes_output.md"

        # Act: Invoke CLI with notes-only
        exit_code = main(
            [
                "--notes-source",
                "file",
                "--notes-file",
                str(notes_file),
                "--notes-label",
                "Meeting Notes – 2025-01-10",
                "--format",
                "md",
                "--output-name",
                "notes_output.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()

        content = output_file.read_text(encoding="utf-8")
        assert "Meeting Notes – 2025-01-10" in content
        assert "Action item: Review Q4 results" in content
        assert "Decision: Approve budget increase" in content

    def test_notes_only_docx(self, tmp_path: Path) -> None:
        """Test CLI creating DOCX document with notes only.

        Exercises: transform_notes(), format_to_docx()
        """
        # Arrange
        notes_file = tmp_path / "notes.txt"
        notes_file.write_text(
            "Summary paragraph with important details.\n\n"
            "* Key point one\n"
            "* Key point two\n"
            "* Key point three\n",
            encoding="utf-8",
        )

        output_file = tmp_path / "notes_output.docx"

        # Act
        exit_code = main(
            [
                "--notes-source",
                "file",
                "--notes-file",
                str(notes_file),
                "--notes-label",
                "Project Notes",
                "--format",
                "docx",
                "--output-name",
                "notes_output.docx",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()

        doc = DocxDocument(str(output_file))  # type: ignore[no-untyped-call]
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Project Notes" in all_text
        assert "Key point one" in all_text

    def test_notes_only_rtf(self, tmp_path: Path) -> None:
        """Test CLI creating RTF document with notes only.

        Exercises: transform_notes(), format_to_rtf()
        """
        # Arrange
        notes_file = tmp_path / "notes.txt"
        notes_file.write_text(
            "- Important decision made\n" "- Follow-up required\n", encoding="utf-8"
        )

        output_file = tmp_path / "notes_output.rtf"

        # Act
        exit_code = main(
            [
                "--notes-source",
                "file",
                "--notes-file",
                str(notes_file),
                "--notes-label",
                "Meeting Summary",
                "--format",
                "rtf",
                "--output-name",
                "notes_output.rtf",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()

        content = output_file.read_text(encoding="utf-8")
        assert r"{\rtf" in content


class TestCLINotesAndTranscript:
    """Test CLI with combined notes and transcript workflows."""

    def test_notes_and_transcript_markdown(self, tmp_path: Path) -> None:
        """Test CLI creating document with both notes and transcript.

        Exercises: Full pipeline with both notes and transcript processing.
        """
        # Arrange: Create notes file
        notes_file = tmp_path / "notes.md"
        notes_file.write_text(
            "- Discussed SpaceX launch\n" "- Team excited about landing\n",
            encoding="utf-8",
        )

        # Create transcript file
        transcript_file = tmp_path / "transcript.txt"
        transcript_file.write_text(SPACEX_DISCUSSION.input_text, encoding="utf-8")

        output_file = tmp_path / "combined_output.md"

        # Act
        exit_code = main(
            [
                "--source",
                "file",
                "--file",
                str(transcript_file),
                "--num-speakers",
                "3",
                "--notes-source",
                "file",
                "--notes-file",
                str(notes_file),
                "--notes-label",
                "Discussion Notes",
                "--format",
                "md",
                "--output-name",
                "combined_output.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()

        content = output_file.read_text(encoding="utf-8")
        # Should have both notes and transcript
        assert "Discussion Notes" in content
        assert "Discussed SpaceX launch" in content
        assert "Speaker A:" in content or "Speaker B:" in content

    def test_notes_and_transcript_docx(self, tmp_path: Path) -> None:
        """Test CLI creating DOCX with both notes and transcript.

        Uses speakerless transcript to avoid UI prompts.
        """
        # Arrange
        notes_file = tmp_path / "notes.txt"
        notes_file.write_text("- Meeting overview\n", encoding="utf-8")

        transcript_file = tmp_path / "transcript.txt"
        # Use speakerless transcript
        transcript_file.write_text(
            "Welcome everyone. Thank you for joining. Let's get started.\n",
            encoding="utf-8",
        )

        output_file = tmp_path / "combined_output.docx"

        # Act
        exit_code = main(
            [
                "--source",
                "file",
                "--file",
                str(transcript_file),
                "--num-speakers",
                "2",
                "--notes-source",
                "file",
                "--notes-file",
                str(notes_file),
                "--notes-label",
                "Quick Notes",
                "--format",
                "docx",
                "--output-name",
                "combined_output.docx",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()

        doc = DocxDocument(str(output_file))  # type: ignore[no-untyped-call]
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Quick Notes" in all_text or "Meeting overview" in all_text
        assert "Speaker" in all_text


class TestCLIUpdateModeNotes:
    """Test CLI update mode for adding/replacing notes.

    These tests exercise document/reader.py via --mode update.
    """

    def test_add_notes_to_existing_markdown(self, tmp_path: Path) -> None:
        """Test CLI adding notes to existing Markdown document.

        Exercises: cli.main() with --mode update, read_document(), merge_notes()
        """
        # Arrange: Create initial document with notes only (no UI prompts)
        initial_notes_file = tmp_path / "initial_notes.txt"
        initial_notes_file.write_text("- Initial note\n", encoding="utf-8")

        initial_output = tmp_path / "initial.md"

        # Create initial document with notes only
        main(
            [
                "--notes-source",
                "file",
                "--notes-file",
                str(initial_notes_file),
                "--notes-label",
                "Initial Notes",
                "--format",
                "md",
                "--output-name",
                "initial.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Create notes to add
        notes_file = tmp_path / "notes.md"
        notes_file.write_text("- Added note\n", encoding="utf-8")

        final_output = tmp_path / "final.md"

        # Act: Add notes to existing document
        exit_code = main(
            [
                "--mode",
                "update",
                "--update-file",
                str(initial_output),
                "--update-action",
                "add-notes",
                "--notes-source",
                "file",
                "--notes-file",
                str(notes_file),
                "--notes-label",
                "Additional Notes",
                "--format",
                "md",
                "--output-name",
                "final.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert final_output.exists()

        content = final_output.read_text(encoding="utf-8")
        assert "Additional Notes" in content
        assert "Added note" in content
        # Original notes should still be present
        assert "Initial Notes" in content or "Initial note" in content

    def test_replace_notes_in_existing_docx(self, tmp_path: Path) -> None:
        """Test CLI replacing notes in existing DOCX document.

        Exercises: read_document() with DOCX, merge_notes() with replace mode
        """
        # Arrange: Create initial document with notes
        initial_notes = tmp_path / "old_notes.txt"
        initial_notes.write_text("- Old note content\n", encoding="utf-8")

        initial_output = tmp_path / "initial.docx"

        main(
            [
                "--notes-source",
                "file",
                "--notes-file",
                str(initial_notes),
                "--notes-label",
                "Old Notes",
                "--format",
                "docx",
                "--output-name",
                "initial.docx",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Create new notes
        new_notes = tmp_path / "new_notes.txt"
        new_notes.write_text("- New note content\n", encoding="utf-8")

        final_output = tmp_path / "final.docx"

        # Act: Replace notes
        exit_code = main(
            [
                "--mode",
                "update",
                "--update-file",
                str(initial_output),
                "--update-action",
                "replace-notes",
                "--notes-source",
                "file",
                "--notes-file",
                str(new_notes),
                "--notes-label",
                "New Notes",
                "--format",
                "docx",
                "--output-name",
                "final.docx",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert final_output.exists()

        doc = DocxDocument(str(final_output))  # type: ignore[no-untyped-call]
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "New Notes" in all_text or "New note content" in all_text


class TestCLIUpdateModeTranscript:
    """Test CLI update mode for adding/replacing transcript.

    Exercises document/reader.py and merge_transcript().
    """

    def test_add_transcript_to_existing_notes(self, tmp_path: Path) -> None:
        """Test CLI adding transcript to document with notes only.

        Exercises: read_document(), merge_transcript() with add mode
        Uses speakerless transcript to avoid UI prompts.
        """
        # Arrange: Create initial document with notes only
        notes_file = tmp_path / "notes.txt"
        notes_file.write_text("- Meeting notes\n", encoding="utf-8")

        initial_output = tmp_path / "initial.md"

        main(
            [
                "--notes-source",
                "file",
                "--notes-file",
                str(notes_file),
                "--notes-label",
                "Notes",
                "--format",
                "md",
                "--output-name",
                "initial.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Create speakerless transcript to add (avoids UI prompts)
        transcript_file = tmp_path / "transcript.txt"
        transcript_file.write_text(
            "New transcript content here. Added later for testing.\n", encoding="utf-8"
        )

        final_output = tmp_path / "final.md"

        # Act: Add transcript with num-speakers for speakerless detection
        exit_code = main(
            [
                "--mode",
                "update",
                "--update-file",
                str(initial_output),
                "--update-action",
                "add-transcript",
                "--source",
                "file",
                "--file",
                str(transcript_file),
                "--num-speakers",
                "2",
                "--format",
                "md",
                "--output-name",
                "final.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert final_output.exists()

        content = final_output.read_text(encoding="utf-8")
        assert "Notes" in content or "Meeting notes" in content
        assert "Speaker" in content or "transcript content" in content

    def test_replace_transcript_in_existing_document(self, tmp_path: Path) -> None:
        """Test CLI replacing transcript in existing document.

        Exercises: read_document(), merge_transcript() with replace mode
        """
        # Arrange: Create initial document with transcript
        old_transcript = tmp_path / "old_transcript.txt"
        old_transcript.write_text("Speaker A: Old content.\n", encoding="utf-8")

        initial_output = tmp_path / "initial.md"

        main(
            [
                "--source",
                "file",
                "--file",
                str(old_transcript),
                "--format",
                "md",
                "--output-name",
                "initial.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Create new transcript
        new_transcript = tmp_path / "new_transcript.txt"
        new_transcript.write_text("Speaker B: New content.\n", encoding="utf-8")

        final_output = tmp_path / "final.md"

        # Act: Replace transcript
        exit_code = main(
            [
                "--mode",
                "update",
                "--update-file",
                str(initial_output),
                "--update-action",
                "replace-transcript",
                "--source",
                "file",
                "--file",
                str(new_transcript),
                "--format",
                "md",
                "--output-name",
                "final.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert final_output.exists()

        content = final_output.read_text(encoding="utf-8")
        assert "Speaker B" in content or "New content" in content


class TestCLIUpdateModeDOCX:
    """Test CLI update mode with DOCX format.

    Exercises document/reader.py with DOCX files.
    """

    def test_add_notes_to_docx_document(self, tmp_path: Path) -> None:
        """Test CLI adding notes to existing DOCX document.

        Exercises: read_document() with DOCX format
        """
        # Arrange: Create initial DOCX document with notes
        initial_notes_file = tmp_path / "initial_notes.txt"
        initial_notes_file.write_text("- Initial content\n", encoding="utf-8")

        initial_output = tmp_path / "initial.docx"

        main(
            [
                "--notes-source",
                "file",
                "--notes-file",
                str(initial_notes_file),
                "--notes-label",
                "Initial Notes",
                "--format",
                "docx",
                "--output-name",
                "initial.docx",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Create notes to add
        notes_file = tmp_path / "notes.txt"
        notes_file.write_text("- DOCX note added\n", encoding="utf-8")

        final_output = tmp_path / "final.docx"

        # Act: Add notes to DOCX
        exit_code = main(
            [
                "--mode",
                "update",
                "--update-file",
                str(initial_output),
                "--update-action",
                "add-notes",
                "--notes-source",
                "file",
                "--notes-file",
                str(notes_file),
                "--notes-label",
                "Additional DOCX Notes",
                "--format",
                "docx",
                "--output-name",
                "final.docx",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert final_output.exists()

        doc = DocxDocument(str(final_output))  # type: ignore[no-untyped-call]
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "DOCX note added" in all_text or "Additional DOCX Notes" in all_text


class TestCLINotesErrorHandling:
    """Test CLI error handling for notes workflows."""

    def test_update_mode_missing_update_file(self, tmp_path: Path) -> None:
        """Test CLI handles missing update file gracefully.

        Exercises error handling in read_document().
        """
        # Arrange
        nonexistent_file = tmp_path / "nonexistent.md"
        notes_file = tmp_path / "notes.txt"
        notes_file.write_text("- Note\n", encoding="utf-8")

        # Act
        exit_code = main(
            [
                "--mode",
                "update",
                "--update-file",
                str(nonexistent_file),
                "--update-action",
                "add-notes",
                "--notes-source",
                "file",
                "--notes-file",
                str(notes_file),
                "--format",
                "md",
                "--output-name",
                "output.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert: Should return non-zero exit code
        assert exit_code != 0

    def test_update_mode_missing_notes_file(self, tmp_path: Path) -> None:
        """Test CLI handles missing notes file gracefully."""
        # Arrange: Create initial document
        initial_file = tmp_path / "initial.md"
        initial_file.write_text("# Initial\n\nContent\n", encoding="utf-8")

        nonexistent_notes = tmp_path / "nonexistent.txt"

        # Act
        exit_code = main(
            [
                "--mode",
                "update",
                "--update-file",
                str(initial_file),
                "--update-action",
                "add-notes",
                "--notes-source",
                "file",
                "--notes-file",
                str(nonexistent_notes),
                "--format",
                "md",
                "--output-name",
                "output.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert: Should return non-zero exit code
        assert exit_code != 0
