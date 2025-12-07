"""End-to-end CLI tests for speakerless transcript workflows.

Tests the CLI interface with speakerless transcripts to verify:
- Argument parsing and validation
- Speakerless detection with multi-speaker fixtures
- Output generation in DOCX/MD/RTF formats
- Coverage of cli.py, extract/from_file.py, and formatters
"""

from pathlib import Path

from docx import Document as DocxDocument  # type: ignore[import-untyped]

from tests.fixtures.multi_speaker import (
    GENERIC_MEETING_3SPEAKER,
    PANEL_DISCUSSION_4SPEAKER,
    SPACEX_DISCUSSION,
    TEAM_STANDUP_3SPEAKER,
)
from transcript_etl_pipeline.cli import main


class TestCLISpeakerless3SpeakersDOCX:
    """Test CLI with 3-speaker speakerless transcripts - DOCX output."""

    def test_spacex_discussion_3speakers_docx(self, tmp_path: Path) -> None:
        """Test CLI with SpaceX 3-speaker discussion producing DOCX output.

        Exercises: cli.main(), argument parsing, extract_from_file(),
        speakerless detection, format_to_docx()
        """
        # Arrange: Create input file with SpaceX fixture
        input_file = tmp_path / "spacex_input.txt"
        input_file.write_text(SPACEX_DISCUSSION.input_text, encoding="utf-8")

        output_file = tmp_path / "spacex_output.docx"

        # Act: Invoke CLI
        exit_code = main(
            [
                "--source",
                "file",
                "--file",
                str(input_file),
                "--num-speakers",
                "3",
                "--format",
                "docx",
                "--output-name",
                "spacex_output.docx",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert: Successful execution and output exists
        assert exit_code == 0
        assert output_file.exists()
        assert output_file.stat().st_size > 0

        # Verify DOCX contains speaker labels
        doc = DocxDocument(str(output_file))  # type: ignore[no-untyped-call]
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Speaker A" in all_text or "Speaker B" in all_text

    def test_generic_meeting_3speakers_docx(self, tmp_path: Path) -> None:
        """Test CLI with generic meeting 3-speaker transcript producing DOCX.

        Uses GENERIC_MEETING_3SPEAKER fixture with formal business dialogue.
        """
        # Arrange
        input_file = tmp_path / "meeting_input.txt"
        input_file.write_text(GENERIC_MEETING_3SPEAKER.input_text, encoding="utf-8")

        output_file = tmp_path / "meeting_output.docx"

        # Act
        exit_code = main(
            [
                "--source",
                "file",
                "--file",
                str(input_file),
                "--num-speakers",
                "3",
                "--format",
                "docx",
                "--output-name",
                "meeting_output.docx",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()

        doc = DocxDocument(str(output_file))  # type: ignore[no-untyped-call]
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Peter Parker" in all_text or "Frank Oz" in all_text

    def test_team_standup_3speakers_docx(self, tmp_path: Path) -> None:
        """Test CLI with team standup 3-speaker transcript producing DOCX.

        Uses TEAM_STANDUP_3SPEAKER fixture with technical standup dialogue.
        """
        # Arrange
        input_file = tmp_path / "standup_input.txt"
        input_file.write_text(TEAM_STANDUP_3SPEAKER.input_text, encoding="utf-8")

        output_file = tmp_path / "standup_output.docx"

        # Act
        exit_code = main(
            [
                "--source",
                "file",
                "--file",
                str(input_file),
                "--num-speakers",
                "3",
                "--format",
                "docx",
                "--output-name",
                "standup_output.docx",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()

        doc = DocxDocument(str(output_file))  # type: ignore[no-untyped-call]
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Speaker" in all_text


class TestCLISpeakerless4SpeakersDOCX:
    """Test CLI with 4-speaker speakerless transcripts - DOCX output."""

    def test_panel_discussion_4speakers_docx(self, tmp_path: Path) -> None:
        """Test CLI with panel discussion 4-speaker transcript producing DOCX.

        Uses PANEL_DISCUSSION_4SPEAKER fixture with AI ethics panel.
        Exercises 4-speaker detection logic.
        """
        # Arrange
        input_file = tmp_path / "panel_input.txt"
        input_file.write_text(PANEL_DISCUSSION_4SPEAKER.input_text, encoding="utf-8")

        output_file = tmp_path / "panel_output.docx"

        # Act
        exit_code = main(
            [
                "--source",
                "file",
                "--file",
                str(input_file),
                "--num-speakers",
                "4",
                "--format",
                "docx",
                "--output-name",
                "panel_output.docx",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()

        doc = DocxDocument(str(output_file))  # type: ignore[no-untyped-call]
        all_text = " ".join(p.text for p in doc.paragraphs)
        # Should have 4 distinct speakers
        assert "Speaker A" in all_text
        assert "Speaker B" in all_text or "Speaker C" in all_text


class TestCLISpeakerlessMarkdown:
    """Test CLI with speakerless transcripts - Markdown output."""

    def test_spacex_discussion_3speakers_md(self, tmp_path: Path) -> None:
        """Test CLI with SpaceX 3-speaker discussion producing Markdown output.

        Exercises: cli.main(), format_to_md()
        """
        # Arrange
        input_file = tmp_path / "spacex_input.txt"
        input_file.write_text(SPACEX_DISCUSSION.input_text, encoding="utf-8")

        output_file = tmp_path / "spacex_output.md"

        # Act
        exit_code = main(
            [
                "--source",
                "file",
                "--file",
                str(input_file),
                "--num-speakers",
                "3",
                "--format",
                "md",
                "--output-name",
                "spacex_output.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()

        content = output_file.read_text(encoding="utf-8")
        assert "Speaker A:" in content or "Speaker B:" in content
        # Check for actual dialogue content
        assert "SpaceX" in content or "launch" in content

    def test_generic_meeting_3speakers_md(self, tmp_path: Path) -> None:
        """Test CLI with generic meeting producing Markdown output."""
        # Arrange
        input_file = tmp_path / "meeting_input.txt"
        input_file.write_text(GENERIC_MEETING_3SPEAKER.input_text, encoding="utf-8")

        output_file = tmp_path / "meeting_output.md"

        # Act
        exit_code = main(
            [
                "--source",
                "file",
                "--file",
                str(input_file),
                "--num-speakers",
                "3",
                "--format",
                "md",
                "--output-name",
                "meeting_output.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()

        content = output_file.read_text(encoding="utf-8")
        assert "Peter Parker" in content or "Frank Oz" in content


class TestCLISpeakerlessRTF:
    """Test CLI with speakerless transcripts - RTF output."""

    def test_spacex_discussion_3speakers_rtf(self, tmp_path: Path) -> None:
        """Test CLI with SpaceX 3-speaker discussion producing RTF output.

        Exercises: cli.main(), format_to_rtf()
        """
        # Arrange
        input_file = tmp_path / "spacex_input.txt"
        input_file.write_text(SPACEX_DISCUSSION.input_text, encoding="utf-8")

        output_file = tmp_path / "spacex_output.rtf"

        # Act
        exit_code = main(
            [
                "--source",
                "file",
                "--file",
                str(input_file),
                "--num-speakers",
                "3",
                "--format",
                "rtf",
                "--output-name",
                "spacex_output.rtf",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()
        assert output_file.stat().st_size > 0

        # RTF format check - should contain RTF header
        content = output_file.read_text(encoding="utf-8")
        assert r"{\rtf" in content

    def test_panel_discussion_4speakers_rtf(self, tmp_path: Path) -> None:
        """Test CLI with panel discussion producing RTF output."""
        # Arrange
        input_file = tmp_path / "panel_input.txt"
        input_file.write_text(PANEL_DISCUSSION_4SPEAKER.input_text, encoding="utf-8")

        output_file = tmp_path / "panel_output.rtf"

        # Act
        exit_code = main(
            [
                "--source",
                "file",
                "--file",
                str(input_file),
                "--num-speakers",
                "4",
                "--format",
                "rtf",
                "--output-name",
                "panel_output.rtf",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()

        content = output_file.read_text(encoding="utf-8")
        assert r"{\rtf" in content


class TestCLISpeakerlessAutoDetect:
    """Test CLI speakerless detection without explicit num-speakers."""

    def test_auto_detect_3speakers(self, tmp_path: Path) -> None:
        """Test CLI auto-detects 3 speakers without --num-speakers flag.

        Verifies that speakerless detection works when num_speakers is not
        explicitly provided (should auto-detect 2-4 speakers).
        """
        # Arrange
        input_file = tmp_path / "spacex_input.txt"
        input_file.write_text(SPACEX_DISCUSSION.input_text, encoding="utf-8")

        output_file = tmp_path / "spacex_auto.md"

        # Act - No --num-speakers flag
        exit_code = main(
            [
                "--source",
                "file",
                "--file",
                str(input_file),
                "--format",
                "md",
                "--output-name",
                "spacex_auto.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert
        assert exit_code == 0
        assert output_file.exists()

        content = output_file.read_text(encoding="utf-8")
        # Should have speaker labels from auto-detection
        assert "Speaker A:" in content or "Speaker B:" in content


class TestCLIErrorHandling:
    """Test CLI error handling for invalid inputs."""

    def test_missing_input_file(self, tmp_path: Path) -> None:
        """Test CLI handles missing input file gracefully.

        Verifies error handling path in cli.main() and extract_from_file().
        """
        # Arrange - File that doesn't exist
        nonexistent_file = tmp_path / "nonexistent.txt"

        # Act
        exit_code = main(
            [
                "--source",
                "file",
                "--file",
                str(nonexistent_file),
                "--format",
                "md",
                "--output-name",
                "output.md",
                "--output-folder",
                str(tmp_path),
            ]
        )

        # Assert - Should return non-zero exit code
        assert exit_code != 0

    def test_invalid_format(self, tmp_path: Path) -> None:
        """Test CLI rejects invalid output format.

        Tests argument validation in create_parser().
        """
        # Arrange
        input_file = tmp_path / "input.txt"
        input_file.write_text("Test content", encoding="utf-8")

        # Act - Invalid format should be caught by argparse
        # This will raise SystemExit, so we catch it
        try:
            main(
                [
                    "--source",
                    "file",
                    "--file",
                    str(input_file),
                    "--format",
                    "invalid",  # Invalid format
                    "--output-name",
                    "output.txt",
                    "--output-folder",
                    str(tmp_path),
                ]
            )
            # If we get here, it didn't raise - fail the test
            raise AssertionError("Should have raised SystemExit for invalid format")
        except SystemExit as e:
            # Assert - Should exit with error code
            assert e.code != 0
