"""DOCX formatter for transcript documents.

This module generates Microsoft Word (.docx) files with proper formatting
according to the transcript formatting rules.
"""

from typing import TYPE_CHECKING, Any

from docx import Document as DocxDocument  # type: ignore[import-untyped]
from docx.enum.text import WD_LINE_SPACING  # type: ignore[import-untyped]
from docx.shared import Pt  # type: ignore[import-untyped]

from transcript_etl_pipeline.document.formatting_rules import (
    BODY_FONT,
    LABEL_FONT,
    SPACING_RULES,
    FontStyle,
    SpacingRule,
)
from transcript_etl_pipeline.document.model import (
    Document,
    DocumentSection,
    Paragraph,
    SectionType,
)

if TYPE_CHECKING:
    # For type checking, we use Any to avoid untyped import issues
    # The actual types from python-docx are not well-typed
    DocxDocType = Any
    DocxParagraphType = Any
    DocxRunType = Any
else:
    DocxDocType = Any
    DocxParagraphType = Any
    DocxRunType = Any

# Mapping from heading level to Word style name
HEADING_STYLES = {
    1: "Heading 1",
    2: "Heading 2",
    3: "Heading 3",
}

# Mapping from bullet level to Word style name
BULLET_STYLES = {
    1: "List Bullet",
    2: "List Bullet 2",
}


def format_to_docx(doc: Document, output_path: str) -> None:
    """Format transcript document to DOCX file.

    Applies all formatting rules:
    - 10pt Calibri font
    - Bold labels, normal body text
    - Single line spacing (1.0)
    - No extra spacing for metadata
    - 12pt spacing above Transcript: label and speaker paragraphs
    - 6pt spacing above regular paragraphs
    - Notes headers use Word heading styles (Heading 1, 2, 3)
    - Notes bullets use Word list styles (List Bullet, List Bullet 2)

    Args:
        doc: The transcript document to format
        output_path: Path where the DOCX file should be saved
    """
    docx_doc = DocxDocument()  # type: ignore[no-untyped-call]

    # Process all sections
    for section in doc.sections:
        _format_section(docx_doc, section)

    # Save the document
    docx_doc.save(output_path)  # type: ignore[no-untyped-call]


def _format_section(docx_doc: Any, section: DocumentSection) -> None:
    """Format a document section.

    Args:
        docx_doc: The python-docx Document object
        section: The transcript section to format
    """
    for paragraph in section.paragraphs:
        _format_paragraph(docx_doc, paragraph, section.section_type)


def _format_paragraph(docx_doc: Any, paragraph: Paragraph, section_type: SectionType) -> None:
    """Format a single paragraph.

    Args:
        docx_doc: The python-docx Document object
        paragraph: The paragraph to format
        section_type: The type of section this paragraph belongs to
    """
    # Handle notes headers with Word heading styles
    if section_type == SectionType.NOTES_HEADER or paragraph.heading_level > 0:
        heading_level = paragraph.heading_level if paragraph.heading_level > 0 else 2
        style_name = HEADING_STYLES.get(heading_level, "Heading 2")
        docx_paragraph = docx_doc.add_paragraph(paragraph.text, style=style_name)
        return

    # Handle notes body with bullets using Word list styles
    if section_type == SectionType.NOTES_BODY:
        if paragraph.is_bullet:
            bullet_level = paragraph.bullet_level if paragraph.bullet_level > 0 else 1
            style_name = BULLET_STYLES.get(bullet_level, "List Bullet")
            docx_paragraph = docx_doc.add_paragraph(paragraph.text, style=style_name)
        else:
            # Non-bullet notes body paragraph
            docx_paragraph = docx_doc.add_paragraph()
            run = docx_paragraph.add_run(paragraph.text)
            _apply_font(run, BODY_FONT)
            spacing = SPACING_RULES[section_type]
            _apply_spacing(docx_paragraph, spacing)
        return

    # Create a new paragraph in the document for transcript/metadata
    docx_paragraph = docx_doc.add_paragraph()

    # Get spacing based on section type
    spacing = SPACING_RULES[section_type]
    _apply_spacing(docx_paragraph, spacing)

    # Add label if present (bold)
    if paragraph.label:
        run = docx_paragraph.add_run(f"{paragraph.label.text} ")
        _apply_font(run, LABEL_FONT)

    # Add body text (normal)
    if paragraph.text:
        run = docx_paragraph.add_run(paragraph.text)
        _apply_font(run, BODY_FONT)


def _apply_spacing(docx_paragraph: Any, spacing: SpacingRule) -> None:
    """Apply spacing rules to a paragraph.

    Args:
        docx_paragraph: The python-docx Paragraph object
        spacing: The spacing rule to apply
    """
    # Set line spacing to 1.0 (single spacing)
    docx_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    # Set before spacing
    if spacing.before_pt > 0:
        docx_paragraph.paragraph_format.space_before = Pt(spacing.before_pt)
    else:
        docx_paragraph.paragraph_format.space_before = Pt(0)

    # Set after spacing
    if spacing.after_pt > 0:
        docx_paragraph.paragraph_format.space_after = Pt(spacing.after_pt)
    else:
        docx_paragraph.paragraph_format.space_after = Pt(0)


def _apply_font(run: Any, font_style: FontStyle) -> None:
    """Apply font styling to a text run.

    Args:
        run: The python-docx Run object
        font_style: The font style to apply
    """
    run.font.name = font_style.name
    run.font.size = Pt(font_style.size_pt)
    run.font.bold = font_style.bold
